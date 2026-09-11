from __future__ import annotations

import io
from pathlib import Path
from typing import Any
from urllib.parse import quote

from app.pipeline.source_normalize import AUDIO_EXTS, VIDEO_EXTS


def utf8_download_stem(title: str) -> str:
    stem = title or "export"
    suffix = Path(stem).suffix.lower()
    if suffix in AUDIO_EXTS | VIDEO_EXTS:
        stem = Path(stem).stem
    cleaned = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in stem)
    return cleaned.strip("._")[:80] or "export"


def ascii_download_stem(title: str) -> str:
    stem = "".join(
        ch if ch.isascii() and (ch.isalnum() or ch in "-_.") else "_"
        for ch in utf8_download_stem(title)
    )
    return stem[:60].strip("._") or "export"


def _ascii_header_filename(name: str) -> str:
    safe = "".join(ch if ch.isascii() and (ch.isalnum() or ch in "-_.") else "_" for ch in name)
    return safe.strip("._")[:80] or "export"


def content_disposition(ascii_name: str, utf8_name: str) -> str:
    """RFC 5987 header: latin-1 `filename=` plus UTF-8 `filename*`."""
    safe_ascii = _ascii_header_filename(ascii_name)
    return (
        f'attachment; filename="{safe_ascii}"; '
        f"filename*=UTF-8''{quote(utf8_name, safe='')}"
    )


def timeline_to_markdown(
    *,
    title: str,
    segments: list[dict[str, Any]],
    summary: str | None = None,
) -> str:
    lines = [f"# {title}", ""]
    if summary:
        lines.extend(["## 摘要", "", summary, "", "## 時間軸", ""])
    else:
        lines.extend(["## 時間軸", ""])
    for seg in segments:
        start = float(seg.get("start") or 0)
        end = float(seg.get("end") or start)
        kind = seg.get("type") or "speech"
        speaker = seg.get("speaker")
        prefix = f"[{start:.1f}–{end:.1f}] ({kind})"
        if speaker:
            prefix += f" {speaker}"
        lines.append(f"- {prefix} {seg.get('text') or ''}")
    lines.append("")
    return "\n".join(lines)


def export_markdown(path: Path, markdown: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(markdown, encoding="utf-8")
    return path


def _docx_bytes(title: str, markdown: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for line in markdown.splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_docx(path: Path, title: str, markdown: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_docx_bytes(title, markdown))
    return path


def _pdf_bytes(title: str, markdown: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font_name = "Helvetica"
    for candidate in (
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        p = Path(candidate)
        if p.exists():
            try:
                pdfmetrics.registerFont(TTFont("ExportFont", str(p)))
                font_name = "ExportFont"
                break
            except Exception:
                continue

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _width, height = A4
    y = height - 48
    c.setFont(font_name, 14)
    c.drawString(48, y, title[:80])
    y -= 28
    c.setFont(font_name, 10)
    for line in markdown.splitlines():
        if y < 48:
            c.showPage()
            c.setFont(font_name, 10)
            y = height - 48
        text = line[:110]
        c.drawString(48, y, text)
        y -= 14
    c.save()
    return buf.getvalue()


def export_pdf(path: Path, title: str, markdown: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_pdf_bytes(title, markdown))
    return path


def build_export_bytes(fmt: str, title: str, markdown: str) -> tuple[bytes, str, str]:
    """Return (bytes, media_type, ascii filename). CJK belongs in filename* only."""
    safe = ascii_download_stem(title)
    if fmt == "md":
        return markdown.encode("utf-8"), "text/markdown; charset=utf-8", f"{safe}.md"
    if fmt == "docx":
        return (
            _docx_bytes(title, markdown),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{safe}.docx",
        )
    if fmt == "pdf":
        return _pdf_bytes(title, markdown), "application/pdf", f"{safe}.pdf"
    raise ValueError(f"Unsupported format: {fmt}")
